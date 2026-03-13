import io
import os
import secrets
import sqlite3
import zipfile
from datetime import datetime
from functools import wraps

import qrcode
from flask import (
    Flask,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DB_PATH = os.path.join(BASE_DIR, 'app.db')

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'change-this-secret-key')
app.config['ADMIN_USERNAME'] = os.getenv('ADMIN_USERNAME', 'admin')
app.config['ADMIN_PASSWORD'] = os.getenv('ADMIN_PASSWORD', 'admin123')
app.config['MAX_CONTENT_LENGTH'] = 10 * 1024 * 1024

CHECK_FIELDS = [
    ('check_production_date', '生产日期'),
    ('check_factory_date', '出厂日期'),
    ('check_team_name', '生成班组'),
    ('check_warehouse', '出厂仓库'),
    ('check_batch_code', '批次编码'),
    ('check_name', '酒类名称'),
    ('check_spec', '酒类规格'),
    ('check_abv', '酒精度数'),
    ('check_dealer', '经销商'),
    ('check_inspection_count', '鉴定次数'),
    ('check_chip', '加密芯片识别'),
    ('check_channel', '销售渠道对比'),
]

IMPORT_HEADERS = [
    ('追溯码', 'trace_code'),
    ('酒类名称', 'name'),
    ('酒类规格', 'spec'),
    ('酒精度数', 'abv'),
    ('生产日期', 'production_date'),
    ('出厂日期', 'factory_date'),
    ('生成班组', 'team_name'),
    ('出厂仓库', 'warehouse'),
    ('批次编码', 'batch_code'),
    ('经销商', 'dealer'),
    ('销售公司', 'sales_company'),
    ('销售渠道', 'sales_channel'),
    ('鉴别报告时间', 'inspection_time'),
    ('鉴别结论', 'inspection_result'),
    ('鉴定次数', 'inspection_count'),
    ('扫码次数', 'scan_count'),
    ('最后扫码时间', 'last_scan_at'),
    ('状态', 'status'),
    ('校验-生产日期', 'check_production_date'),
    ('校验-出厂日期', 'check_factory_date'),
    ('校验-生成班组', 'check_team_name'),
    ('校验-出厂仓库', 'check_warehouse'),
    ('校验-批次编码', 'check_batch_code'),
    ('校验-酒类名称', 'check_name'),
    ('校验-酒类规格', 'check_spec'),
    ('校验-酒精度数', 'check_abv'),
    ('校验-经销商', 'check_dealer'),
    ('校验-鉴定次数', 'check_inspection_count'),
    ('校验-加密芯片识别', 'check_chip'),
    ('校验-销售渠道对比', 'check_channel'),
]

IMPORT_FIELD_MAP = {label.strip().lower(): field for label, field in IMPORT_HEADERS}
IMPORT_FIELD_MAP.update({field.strip().lower(): field for _, field in IMPORT_HEADERS})


def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop('db', None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.execute(
        '''
        CREATE TABLE IF NOT EXISTS liquor (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trace_code TEXT NOT NULL UNIQUE,
            name TEXT NOT NULL,
            spec TEXT,
            abv TEXT,
            production_date TEXT,
            factory_date TEXT,
            team_name TEXT,
            warehouse TEXT,
            batch_code TEXT,
            dealer TEXT,
            sales_company TEXT,
            sales_channel TEXT,
            inspection_time TEXT,
            inspection_result TEXT DEFAULT '已通过研判',
            inspection_count INTEGER DEFAULT 0,
            scan_count INTEGER DEFAULT 0,
            last_scan_at TEXT,
            status TEXT DEFAULT '正常',
            check_production_date INTEGER DEFAULT 1,
            check_factory_date INTEGER DEFAULT 1,
            check_team_name INTEGER DEFAULT 1,
            check_warehouse INTEGER DEFAULT 1,
            check_batch_code INTEGER DEFAULT 1,
            check_name INTEGER DEFAULT 1,
            check_spec INTEGER DEFAULT 1,
            check_abv INTEGER DEFAULT 1,
            check_dealer INTEGER DEFAULT 1,
            check_inspection_count INTEGER DEFAULT 1,
            check_chip INTEGER DEFAULT 1,
            check_channel INTEGER DEFAULT 1,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        '''
    )
    db.commit()
    db.close()


init_db()


def admin_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin_login', next=request.path))
        return view(*args, **kwargs)

    return wrapped_view


def parse_checkbox(name: str) -> int:
    return 1 if request.form.get(name) == 'on' else 0


def now_str() -> str:
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def normalize_header(value) -> str:
    return str(value or '').strip().lower()


def cell_to_text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return str(value).strip()


def safe_int(value, default=0) -> int:
    text = cell_to_text(value)
    if not text:
        return default
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return default


def parse_bool_value(value, default=1) -> int:
    text = cell_to_text(value)
    if not text:
        return int(default)
    text = text.strip().lower()
    truthy = {'1', 'true', 'yes', 'y', 'on', '是', '√', '通过', 'pass', 'ok'}
    falsy = {'0', 'false', 'no', 'n', 'off', '否', '×', '不通过', 'fail'}
    if text in truthy:
        return 1
    if text in falsy:
        return 0
    return int(default)


def generate_trace_code() -> str:
    db = get_db()
    while True:
        code = secrets.token_hex(8).upper()
        exists = db.execute('SELECT 1 FROM liquor WHERE trace_code = ?', (code,)).fetchone()
        if not exists:
            return code


def build_qrcode_bytes(trace_code: str) -> io.BytesIO:
    target_url = url_for('trace_page', trace_code=trace_code, _external=True)
    qr = qrcode.QRCode(box_size=10, border=2)
    qr.add_data(target_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color='black', back_color='white')
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return buf


def get_filtered_liquors(db, q: str):
    if q:
        pattern = f'%{q}%'
        return db.execute(
            '''
            SELECT * FROM liquor
            WHERE trace_code LIKE ?
               OR name LIKE ?
               OR batch_code LIKE ?
               OR dealer LIKE ?
            ORDER BY id DESC
            ''',
            (pattern, pattern, pattern, pattern),
        ).fetchall()
    return db.execute('SELECT * FROM liquor ORDER BY id DESC').fetchall()


@app.context_processor
def inject_globals():
    return {
        'check_fields': CHECK_FIELDS,
    }


@app.route('/')
def index():
    return redirect(url_for('admin_list')) if session.get('admin_logged_in') else redirect(url_for('admin_login'))


@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        if username == app.config['ADMIN_USERNAME'] and password == app.config['ADMIN_PASSWORD']:
            session['admin_logged_in'] = True
            flash('登录成功。', 'success')
            next_url = request.args.get('next') or url_for('admin_list')
            return redirect(next_url)
        flash('用户名或密码错误。', 'danger')
    return render_template('admin_login.html')


@app.route('/admin/logout', methods=['POST'])
@admin_required
def admin_logout():
    session.clear()
    flash('已退出登录。', 'success')
    return redirect(url_for('admin_login'))


@app.route('/admin/liquors')
@admin_required
def admin_list():
    q = request.args.get('q', '').strip()
    db = get_db()
    liquors = get_filtered_liquors(db, q)
    return render_template('admin_list.html', liquors=liquors, q=q)


@app.route('/admin/liquors/new', methods=['GET', 'POST'])
@admin_required
def admin_new():
    if request.method == 'POST':
        db = get_db()
        trace_code = generate_trace_code()
        current_time = now_str()
        data = collect_form_data(trace_code=trace_code, created_at=current_time, updated_at=current_time)
        columns = ', '.join(data.keys())
        placeholders = ', '.join(['?'] * len(data))
        db.execute(f'INSERT INTO liquor ({columns}) VALUES ({placeholders})', tuple(data.values()))
        db.commit()
        flash('酒品已新增，二维码已可访问。', 'success')
        return redirect(url_for('admin_list'))
    return render_template('admin_form.html', liquor=None, action='new')


@app.route('/admin/liquors/<int:liquor_id>/edit', methods=['GET', 'POST'])
@admin_required
def admin_edit(liquor_id: int):
    db = get_db()
    liquor = db.execute('SELECT * FROM liquor WHERE id = ?', (liquor_id,)).fetchone()
    if not liquor:
        abort(404)

    if request.method == 'POST':
        data = collect_form_data(trace_code=liquor['trace_code'], created_at=liquor['created_at'], updated_at=now_str())
        set_clause = ', '.join([f'{key} = ?' for key in data.keys()])
        values = list(data.values()) + [liquor_id]
        db.execute(f'UPDATE liquor SET {set_clause} WHERE id = ?', values)
        db.commit()
        flash('酒品信息已更新。', 'success')
        return redirect(url_for('admin_list'))

    return render_template('admin_form.html', liquor=liquor, action='edit')


@app.route('/admin/liquors/<int:liquor_id>/delete', methods=['POST'])
@admin_required
def admin_delete(liquor_id: int):
    db = get_db()
    db.execute('DELETE FROM liquor WHERE id = ?', (liquor_id,))
    db.commit()
    flash('酒品记录已删除。', 'success')
    return redirect(url_for('admin_list'))


@app.route('/admin/liquors/import', methods=['GET', 'POST'])
@admin_required
def admin_import():
    if request.method == 'POST':
        uploaded_file = request.files.get('excel_file')
        if not uploaded_file or not uploaded_file.filename:
            flash('请先选择一个 Excel 文件。', 'danger')
            return redirect(url_for('admin_import'))
        if not uploaded_file.filename.lower().endswith('.xlsx'):
            flash('当前仅支持 .xlsx 格式。', 'danger')
            return redirect(url_for('admin_import'))

        try:
            workbook = load_workbook(uploaded_file, data_only=True)
            worksheet = workbook.active
        except Exception:
            flash('Excel 文件读取失败，请确认文件未损坏。', 'danger')
            return redirect(url_for('admin_import'))

        header_row = next(worksheet.iter_rows(min_row=1, max_row=1), None)
        if not header_row:
            flash('Excel 文件为空。', 'danger')
            return redirect(url_for('admin_import'))

        headers = [normalize_header(cell.value) for cell in header_row]
        field_indexes = {}
        for idx, header in enumerate(headers):
            mapped = IMPORT_FIELD_MAP.get(header)
            if mapped:
                field_indexes[mapped] = idx

        if 'name' not in field_indexes:
            flash('Excel 表头缺少“酒类名称”列。请先下载导入模板。', 'danger')
            return redirect(url_for('admin_import'))

        db = get_db()
        inserted = 0
        updated = 0
        skipped = 0
        errors = []

        for row_number, row in enumerate(worksheet.iter_rows(min_row=2), start=2):
            values = [cell.value for cell in row]
            if all(cell_to_text(v) == '' for v in values):
                continue

            current_time = now_str()
            try:
                raw = {}
                for field_name, index in field_indexes.items():
                    raw[field_name] = values[index] if index < len(values) else None

                name = cell_to_text(raw.get('name'))
                if not name:
                    skipped += 1
                    errors.append(f'第 {row_number} 行缺少酒类名称，已跳过。')
                    continue

                trace_code = cell_to_text(raw.get('trace_code'))
                existing = None
                if trace_code:
                    existing = db.execute('SELECT * FROM liquor WHERE trace_code = ?', (trace_code,)).fetchone()
                if not trace_code:
                    trace_code = generate_trace_code()

                data = {
                    'trace_code': trace_code,
                    'name': name,
                    'spec': cell_to_text(raw.get('spec')) or (existing['spec'] if existing else ''),
                    'abv': cell_to_text(raw.get('abv')) or (existing['abv'] if existing else ''),
                    'production_date': cell_to_text(raw.get('production_date')) or (existing['production_date'] if existing else ''),
                    'factory_date': cell_to_text(raw.get('factory_date')) or (existing['factory_date'] if existing else ''),
                    'team_name': cell_to_text(raw.get('team_name')) or (existing['team_name'] if existing else ''),
                    'warehouse': cell_to_text(raw.get('warehouse')) or (existing['warehouse'] if existing else ''),
                    'batch_code': cell_to_text(raw.get('batch_code')) or (existing['batch_code'] if existing else ''),
                    'dealer': cell_to_text(raw.get('dealer')) or (existing['dealer'] if existing else ''),
                    'sales_company': cell_to_text(raw.get('sales_company')) or (existing['sales_company'] if existing else ''),
                    'sales_channel': cell_to_text(raw.get('sales_channel')) or (existing['sales_channel'] if existing else ''),
                    'inspection_time': cell_to_text(raw.get('inspection_time')) or (existing['inspection_time'] if existing else current_time),
                    'inspection_result': cell_to_text(raw.get('inspection_result')) or (existing['inspection_result'] if existing else '已通过研判'),
                    'inspection_count': safe_int(raw.get('inspection_count'), existing['inspection_count'] if existing else 0),
                    'scan_count': safe_int(raw.get('scan_count'), existing['scan_count'] if existing else 0),
                    'last_scan_at': cell_to_text(raw.get('last_scan_at')) or (existing['last_scan_at'] if existing else ''),
                    'status': cell_to_text(raw.get('status')) or (existing['status'] if existing else '正常'),
                    'check_production_date': parse_bool_value(raw.get('check_production_date'), existing['check_production_date'] if existing else 1),
                    'check_factory_date': parse_bool_value(raw.get('check_factory_date'), existing['check_factory_date'] if existing else 1),
                    'check_team_name': parse_bool_value(raw.get('check_team_name'), existing['check_team_name'] if existing else 1),
                    'check_warehouse': parse_bool_value(raw.get('check_warehouse'), existing['check_warehouse'] if existing else 1),
                    'check_batch_code': parse_bool_value(raw.get('check_batch_code'), existing['check_batch_code'] if existing else 1),
                    'check_name': parse_bool_value(raw.get('check_name'), existing['check_name'] if existing else 1),
                    'check_spec': parse_bool_value(raw.get('check_spec'), existing['check_spec'] if existing else 1),
                    'check_abv': parse_bool_value(raw.get('check_abv'), existing['check_abv'] if existing else 1),
                    'check_dealer': parse_bool_value(raw.get('check_dealer'), existing['check_dealer'] if existing else 1),
                    'check_inspection_count': parse_bool_value(raw.get('check_inspection_count'), existing['check_inspection_count'] if existing else 1),
                    'check_chip': parse_bool_value(raw.get('check_chip'), existing['check_chip'] if existing else 1),
                    'check_channel': parse_bool_value(raw.get('check_channel'), existing['check_channel'] if existing else 1),
                    'created_at': existing['created_at'] if existing else current_time,
                    'updated_at': current_time,
                }

                if existing:
                    set_clause = ', '.join([f'{key} = ?' for key in data.keys()])
                    db.execute(f'UPDATE liquor SET {set_clause} WHERE id = ?', list(data.values()) + [existing['id']])
                    updated += 1
                else:
                    columns = ', '.join(data.keys())
                    placeholders = ', '.join(['?'] * len(data))
                    db.execute(f'INSERT INTO liquor ({columns}) VALUES ({placeholders})', tuple(data.values()))
                    inserted += 1
            except Exception as exc:
                skipped += 1
                errors.append(f'第 {row_number} 行导入失败：{exc}')

        db.commit()
        flash(f'批量导入完成。新增 {inserted} 条，更新 {updated} 条，跳过 {skipped} 条。', 'success')
        for message in errors[:5]:
            flash(message, 'warning')
        if len(errors) > 5:
            flash(f'其余 {len(errors) - 5} 条错误未展开显示。', 'warning')
        return redirect(url_for('admin_list'))

    return render_template('admin_import.html', import_headers=IMPORT_HEADERS)


@app.route('/admin/liquors/import-template.xlsx')
@admin_required
def admin_import_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '酒品导入模板'
    headers = [label for label, _ in IMPORT_HEADERS]
    worksheet.append(headers)
    worksheet.append([
        '',
        '新飞天53%vol 500ml贵州茅台酒（1×6）RFID',
        '500ml',
        '53',
        '2020-06-09',
        '2020-06-18 20:48:54',
        '包装车间二班组',
        '贵阳三库',
        '2019-148',
        '廊坊中糖华洋实业有限公司',
        '贵阳三库（京东）',
        '京东',
        '2022-06-09 16:34:38',
        '已通过研判',
        5,
        0,
        '',
        '正常',
        1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1,
    ])

    header_fill = PatternFill('solid', fgColor='D32F2F')
    header_font = Font(color='FFFFFF', bold=True)
    for col_idx, label in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        worksheet.column_dimensions[cell.column_letter].width = max(14, len(label) + 4)

    worksheet.freeze_panes = 'A2'

    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='酒品导入模板.xlsx',
    )


@app.route('/admin/qrcode/<trace_code>/download')
@admin_required
def admin_download_qrcode(trace_code: str):
    db = get_db()
    liquor = db.execute('SELECT trace_code FROM liquor WHERE trace_code = ?', (trace_code,)).fetchone()
    if not liquor:
        abort(404)
    buf = build_qrcode_bytes(trace_code)
    return send_file(buf, mimetype='image/png', as_attachment=True, download_name=f'{trace_code}.png')


@app.route('/admin/qrcodes/download')
@admin_required
def admin_download_qrcodes_zip():
    q = request.args.get('q', '').strip()
    db = get_db()
    liquors = get_filtered_liquors(db, q)
    if not liquors:
        flash('当前没有可下载的二维码。', 'warning')
        return redirect(url_for('admin_list', q=q))

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for liquor in liquors:
            qr_bytes = build_qrcode_bytes(liquor['trace_code'])
            zf.writestr(f"{liquor['trace_code']}.png", qr_bytes.getvalue())
    zip_buf.seek(0)
    filename = '酒品二维码.zip' if not q else f'酒品二维码_{q}.zip'
    return send_file(zip_buf, mimetype='application/zip', as_attachment=True, download_name=filename)


@app.route('/trace/<trace_code>')
def trace_page(trace_code: str):
    db = get_db()
    liquor = db.execute('SELECT * FROM liquor WHERE trace_code = ?', (trace_code,)).fetchone()
    if not liquor:
        abort(404)
    current_time = now_str()
    db.execute(
        'UPDATE liquor SET scan_count = COALESCE(scan_count, 0) + 1, last_scan_at = ?, updated_at = ? WHERE id = ?',
        (current_time, current_time, liquor['id']),
    )
    db.commit()
    liquor = db.execute('SELECT * FROM liquor WHERE trace_code = ?', (trace_code,)).fetchone()
    return render_template('trace.html', liquor=liquor)


@app.route('/qrcode/<trace_code>.png')
def qrcode_image(trace_code: str):
    db = get_db()
    liquor = db.execute('SELECT trace_code FROM liquor WHERE trace_code = ?', (trace_code,)).fetchone()
    if not liquor:
        abort(404)
    buf = build_qrcode_bytes(trace_code)
    return send_file(buf, mimetype='image/png')


@app.route('/seed-demo')
def seed_demo():
    db = get_db()
    existing = db.execute('SELECT COUNT(*) AS cnt FROM liquor').fetchone()['cnt']
    if existing > 0:
        return '已有数据，未重复插入。'

    current_time = now_str()
    trace_code = generate_trace_code()
    db.execute(
        '''
        INSERT INTO liquor (
            trace_code, name, spec, abv, production_date, factory_date,
            team_name, warehouse, batch_code, dealer, sales_company,
            sales_channel, inspection_time, inspection_result,
            inspection_count, scan_count, last_scan_at, status,
            check_production_date, check_factory_date, check_team_name,
            check_warehouse, check_batch_code, check_name, check_spec,
            check_abv, check_dealer, check_inspection_count,
            check_chip, check_channel, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''',
        (
            trace_code, '新飞天53%vol 500ml贵州茅台酒（1×6）RFID', '500ml', '53',
            '2020-06-09', '2020-06-18 20:48:54', '包装车间二班组', '贵阳三库',
            '2019-148', '廊坊中糖华洋实业有限公司', '贵阳三库（京东）', '京东',
            '2022-06-09 16:34:38', '已通过研判', 5, 0, '', '正常',
            1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, current_time, current_time
        ),
    )
    db.commit()
    return f'演示数据已插入，追溯码：{trace_code}'


def collect_form_data(trace_code: str, created_at: str, updated_at: str):
    return {
        'trace_code': trace_code,
        'name': request.form.get('name', '').strip(),
        'spec': request.form.get('spec', '').strip(),
        'abv': request.form.get('abv', '').strip(),
        'production_date': request.form.get('production_date', '').strip(),
        'factory_date': request.form.get('factory_date', '').strip(),
        'team_name': request.form.get('team_name', '').strip(),
        'warehouse': request.form.get('warehouse', '').strip(),
        'batch_code': request.form.get('batch_code', '').strip(),
        'dealer': request.form.get('dealer', '').strip(),
        'sales_company': request.form.get('sales_company', '').strip(),
        'sales_channel': request.form.get('sales_channel', '').strip(),
        'inspection_time': request.form.get('inspection_time', '').strip() or now_str(),
        'inspection_result': request.form.get('inspection_result', '已通过研判').strip() or '已通过研判',
        'inspection_count': int(request.form.get('inspection_count', '0') or 0),
        'scan_count': int(request.form.get('scan_count', '0') or 0),
        'last_scan_at': request.form.get('last_scan_at', '').strip(),
        'status': request.form.get('status', '正常').strip() or '正常',
        'check_production_date': parse_checkbox('check_production_date'),
        'check_factory_date': parse_checkbox('check_factory_date'),
        'check_team_name': parse_checkbox('check_team_name'),
        'check_warehouse': parse_checkbox('check_warehouse'),
        'check_batch_code': parse_checkbox('check_batch_code'),
        'check_name': parse_checkbox('check_name'),
        'check_spec': parse_checkbox('check_spec'),
        'check_abv': parse_checkbox('check_abv'),
        'check_dealer': parse_checkbox('check_dealer'),
        'check_inspection_count': parse_checkbox('check_inspection_count'),
        'check_chip': parse_checkbox('check_chip'),
        'check_channel': parse_checkbox('check_channel'),
        'created_at': created_at,
        'updated_at': updated_at,
    }


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
